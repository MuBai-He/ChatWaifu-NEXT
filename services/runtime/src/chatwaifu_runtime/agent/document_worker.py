"""Trusted document worker CLI. Only declarative content is accepted."""

import json
import sys
from collections.abc import Iterable
from typing import Protocol, cast


class _FontElement(Protocol):
    attrib: dict[str, str]

    def set(self, key: str, value: str) -> None: ...


class _StyleElement(Protocol):
    def iter(self, tag: str) -> Iterable[_FontElement]: ...


def main() -> None:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt
    from docx.styles.style import ParagraphStyle

    data = json.loads(sys.stdin.buffer.read(131073))
    document = Document()
    style = cast(ParagraphStyle, document.styles["Normal"])
    font = data.get("_font_face", "Noto Sans CJK SC")
    for fonts in cast(_StyleElement, document.styles.element).iter(qn("w:rFonts")):
        for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn("w:" + key), None)
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + key), font)
    style.font.name = font
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)  # pyright: ignore[reportUnknownMemberType]
    style.font.size = Pt(11)
    for name in ("Title", "Heading 1"):
        heading = cast(ParagraphStyle, document.styles[name])
        heading.font.name = font
        heading.font.bold = False
        heading.font.cs_bold = False
        heading.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)  # pyright: ignore[reportUnknownMemberType]
    document.add_heading(data["title"], 0)
    for section in data["sections"]:
        if section.get("heading"):
            document.add_heading(section["heading"], 1)
        for paragraph in section["paragraphs"]:
            document.add_paragraph(paragraph)
    document.save(sys.argv[1])


if __name__ == "__main__":
    main()
